from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph
from reportlab.lib.styles import getSampleStyleSheet

doc = SimpleDocTemplate('test_sample.pdf', pagesize=letter)
styles = getSampleStyleSheet()
doc.build([
    Paragraph('Transformers in Natural Language Processing', styles['Heading1']),
    Paragraph(
        'Transformers have revolutionized NLP by introducing self-attention mechanisms. '
        'BERT and GPT are widely used transformer-based models that achieve state-of-the-art '
        'results on tasks like text classification, question answering, and summarization. '
        'Pre-training on large corpora then fine-tuning on specific tasks is the dominant '
        'paradigm in modern NLP research and has proven extremely effective.',
        styles['BodyText']
    ),
])
print('PDF created: test_sample.pdf')

from docx import Document
d = Document()
d.add_heading('Deep Learning Overview', 0)
d.add_paragraph(
    'Deep learning is a subset of machine learning that uses neural networks with multiple '
    'layers. Convolutional neural networks excel at image recognition tasks. Recurrent neural '
    'networks are effective for sequential data like time series and text. Transformers now '
    'dominate NLP tasks thanks to their attention mechanisms and scalability.'
)
d.save('test_sample.docx')
print('DOCX created: test_sample.docx')
